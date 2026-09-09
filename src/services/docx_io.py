"""DOCX document I/O: reading, marker detection, replacement, and bibliography."""

import copy
import os
import re
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import NamedTuple, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.enum.style import WD_STYLE_TYPE

from ..models.markers import MarkerConfig
from ..utils.markers import find_markers as _find_marker_specs
from .docx_fields import (
    FieldIndex, build_field_runs, make_result_run, run_ancestry,
)

logger = logging.getLogger(__name__)

# The (REF)/(REFS)-only pattern; ``find_markers`` uses the full grammar of
# ``src.utils.markers`` (author-suggested citations included).
MARKER_PATTERN = re.compile(r'\((REFS?)\)')

BIBLIOGRAPHY_HEADING_STYLE = 'AIREFS Bibliography Heading'
BIBLIOGRAPHY_ENTRY_STYLE = 'AIREFS Bibliography'
MAX_HEADING_LENGTH = 80

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
MC_NS = 'http://schemas.openxmlformats.org/markup-compatibility/2006'


def _child_text(child) -> str:
    """Text one child of a ``w:r`` contributes to ``paragraph.text``, by
    python-docx's ``CT_R.text`` rules.

    '' for everything that renders without text -- page and column breaks,
    ``w:sym``, drawings, footnote references, ``w:lastRenderedPageBreak`` --
    and for field code (``w:instrText``) and deleted text (``w:delText``).
    """
    tag = child.tag
    if tag == f'{{{W_NS}}}t':
        return child.text or ''
    if tag in (f'{{{W_NS}}}tab', f'{{{W_NS}}}ptab'):
        return '\t'
    if tag == f'{{{W_NS}}}br':
        # Only text-wrapping line breaks are text; page and column
        # breaks are invisible to paragraph.text.
        if child.get(f'{{{W_NS}}}type', 'textWrapping') == 'textWrapping':
            return '\n'
        return ''
    if tag == f'{{{W_NS}}}cr':
        return '\n'
    if tag == f'{{{W_NS}}}noBreakHyphen':
        return '-'
    return ''


def run_text(r_elem) -> str:
    """Visible text of a ``w:r`` element, matching python-docx's ``CT_R.text``.

    Field code (``w:instrText``) and deleted text (``w:delText``) contribute
    nothing, exactly as in ``paragraph.text``. Works on any lxml ``w:r``,
    including nested ones that python-docx's ``Run`` wrapper never exposes.
    """
    return ''.join(_child_text(child) for child in r_elem)


def iter_text_runs(paragraph) -> list:
    """Runs that contribute to ``paragraph.text``, in order.

    python-docx computes ``paragraph.text`` from ``w:r | w:hyperlink`` children,
    so this is exactly ``./w:r | ./w:hyperlink/w:r``. Use it (and only it) for
    character-offset arithmetic; a test asserts equality with paragraph.text.
    """
    return paragraph._p.xpath('./w:r | ./w:hyperlink/w:r')


@dataclass
class RunInfo:
    elem: object            # the w:r element
    deleted: bool           # under w:del / w:moveFrom: gone in the final view
    inserted: bool          # under w:ins / w:moveTo: present in the final view


def iter_all_runs(paragraph):
    """Every ``w:r`` under the paragraph except ``mc:Fallback`` duplicates.

    Reaches runs nested in hyperlinks, tracked changes, content controls and
    text boxes. For field awareness only -- never for character offsets.

    ``deleted`` / ``inserted`` follow final-view Track Changes semantics and
    come from ``docx_fields.run_ancestry``, the same helper the field reader
    uses, so a moved-from run (which keeps ``w:t``, unlike ``w:delText``) is
    flagged deleted here exactly as its field is in ``FieldIndex``.
    """
    p_elem = paragraph._p
    for r in p_elem.iter(f'{{{W_NS}}}r'):
        a = run_ancestry(r, p_elem)
        if not a.skip:
            yield RunInfo(elem=r, deleted=a.deleted, inserted=a.inserted)


class FieldBoundaryError(RuntimeError):
    """Raised when an edit would split or rewrite runs that belong to a field."""


class MarkerSpan(NamedTuple):
    """Where a marker sits among a paragraph's text runs (``_locate_span``)."""
    runs: list          # iter_text_runs(paragraph)
    first: int          # index in runs of the first run the marker touches
    first_off: int      # marker start as an offset into runs[first]
    last: int           # index in runs of the last run the marker touches
    last_off: int       # marker end as an offset into runs[last]


class DocxHandler:
    """Read and manipulate DOCX documents for reference insertion."""

    def __init__(self, path: str):
        self.path = path
        self.doc = Document(path)
        self._fields: Optional[FieldIndex] = None
        # Element that followed the last removed References section; a
        # rebuilt bibliography is inserted before it so it stays in place.
        self._insert_anchor = None

    @property
    def fields(self) -> FieldIndex:
        """Field membership for the current document tree, built on demand.

        Keyed on run element identity, so every method that adds or removes
        a field run calls :meth:`invalidate_fields` afterwards. Building it
        walks the whole body, so an edit that provably leaves every field
        run in place (``_split_and_emit``) keeps the index instead.
        """
        if self._fields is None:
            self._fields = FieldIndex(self.doc)
        return self._fields

    def invalidate_fields(self):
        """Forget the cached :attr:`fields`; call after any structural edit."""
        self._fields = None

    def get_paragraphs(self) -> list:
        """Return all paragraphs in the document."""
        return self.doc.paragraphs

    def get_full_text(self) -> str:
        """Return full document text."""
        return "\n".join(p.text for p in self.doc.paragraphs)

    def find_markers(self, config: Optional[MarkerConfig] = None) -> list[dict]:
        """Find all citation markers in the document.

        ``config`` selects which marker kinds are recognised ((REF)/(REFS)
        only, or author-suggested citations as well); it must match the
        configuration the pipeline ran with so that export replaces exactly
        the markers that were processed (``ProjectState.export_marker_config``).

        Returns list of dicts with keys:
            paragraph, para_index, location, marker_type, text, spec,
            full_paragraph_text, marker_order_in_para.
        ``location`` is the marker's span in ``paragraph.text`` and ``text``
        the marker exactly as written, e.g. ``(REF)`` or ``(Smith et al. 2020)``.
        """
        config = config or MarkerConfig.all_on()
        markers = []
        for para_idx, para in enumerate(self.doc.paragraphs):
            text = para.text
            for order, spec in enumerate(_find_marker_specs(text, config)):
                markers.append({
                    'paragraph': para,
                    'para_index': para_idx,
                    'location': (spec.start, spec.end),
                    'marker_type': spec.kind.value,
                    'text': spec.text,
                    'spec': spec,
                    'full_paragraph_text': text,
                    'marker_order_in_para': order,
                })
        logger.info(f"Found {len(markers)} markers in document")
        return markers

    def replace_marker_by_regex(self, paragraph, marker_text: str, replacement: str,
                                superscript: Optional[bool] = False, start: int = 0):
        """Replace the first occurrence of *marker_text* at or after *start* in *paragraph*.

        Locates the marker over the text runs, then replaces the touched runs
        with up to three new ones (before | citation | after). The before- and
        after-runs keep their neighbours' ``w:rPr`` and every content child
        outside the marker (tabs, breaks, symbols, drawings... included, as
        the elements they are); the citation run is one ``w:t`` under the
        first touched run's ``w:rPr``. Every other run is untouched.

        *superscript* is tri-state. True: the citation run is superscript and
        the after-run loses any superscript (a fresh superscript citation).
        False: the citation run carries no vertical alignment (a fresh
        bracket citation). None: the citation run keeps the marker run's
        vertical alignment -- for rewriting an existing citation in place,
        where a superscript ``[3]`` must stay superscript.

        Returns the new citation ``w:r`` element, or None when the marker is
        not in the paragraph. Raises :class:`FieldBoundaryError` rather than
        rewriting a run that belongs to a field. *start* (an offset into
        ``paragraph.text``) picks a later occurrence when the same marker text
        appears several times; callers writing several markers of one
        paragraph go right to left so earlier offsets stay valid.
        """
        span = self._locate_span(paragraph, marker_text, start)
        if span is None:
            return None
        return self._split_and_emit(span, replacement, superscript)

    def _locate_span(self, paragraph, marker_text: str, start: int = 0) -> Optional[MarkerSpan]:
        """Where the first *marker_text* at or after *start* sits among
        ``iter_text_runs(paragraph)``.

        Offsets are computed over the very runs that make up ``paragraph.text``,
        so a marker present in the text is always found. Raises
        :class:`FieldBoundaryError` if any touched run belongs to a field: a
        cached field result must be rewritten through the field API instead.
        """
        full_text = paragraph.text
        start = full_text.find(marker_text, start)
        if start < 0:
            return None
        end = start + len(marker_text)
        runs = iter_text_runs(paragraph)
        first = last = None
        offset = 0
        for i, r in enumerate(runs):
            run_start, run_end = offset, offset + len(run_text(r))
            if run_start < end and run_end > start:
                if first is None:
                    first = (i, start - run_start)
                last = (i, end - run_start)
            offset = run_end
        if first is None:               # only an empty marker_text overlaps no run
            return None
        span = MarkerSpan(runs, first[0], first[1], last[0], last[1])
        for r in runs[span.first:span.last + 1]:
            if self.fields.in_field(r):
                raise FieldBoundaryError(
                    f"marker {marker_text!r} overlaps a field; refusing to rewrite its runs")
        return span

    @staticmethod
    def _run_with_rpr(template_r, superscript: Optional[bool]):
        """New, otherwise empty ``w:r`` carrying a copy of *template_r*'s
        ``w:rPr`` -- the one child that formatting lives in.

        *superscript*: None keeps the template's vertical alignment, True makes
        the run superscript, False removes any vertical alignment.
        """
        new_r = OxmlElement('w:r')
        rpr = template_r.find(qn('w:rPr'))
        if rpr is not None:
            rpr = copy.deepcopy(rpr)
            new_r.append(rpr)
        if superscript is not None:
            if rpr is not None:
                for va in rpr.findall(qn('w:vertAlign')):
                    rpr.remove(va)
            if superscript:
                if rpr is None:
                    rpr = new_r.get_or_add_rPr()
                rpr.get_or_add_vertAlign().set(qn('w:val'), 'superscript')
        return new_r

    @staticmethod
    def _text_elem(text: str):
        """A ``w:t`` with *text*, whitespace preserved."""
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        return t

    @staticmethod
    def _slice_children(r_elem, lo: int, hi: Optional[int]) -> list:
        """Deep copies of the content children of *r_elem* (everything but
        ``w:rPr``) that lie within characters [lo, hi] of ``run_text(r_elem)``;
        *hi* None means the end of the run.

        A ``w:t`` straddling a bound is cut to the part inside; a tab, break
        or hyphen is one character and so wholly in or out. Children that
        contribute no text (page breaks, ``w:sym``, drawings, footnote
        references...) are kept when they sit anywhere in the range, bounds
        included: a page break right after a marker stays with the text after
        it, a symbol right before stays with the text before.
        """
        out, pos = [], 0
        for child in r_elem:
            if child.tag == qn('w:rPr'):
                continue
            start = pos
            end = pos = start + len(_child_text(child))
            if start == end:                                    # zero-width
                if lo <= start and (hi is None or start <= hi):
                    out.append(copy.deepcopy(child))
                continue
            cut_lo = max(start, lo)
            cut_hi = end if hi is None else min(end, hi)
            if cut_lo >= cut_hi:
                continue
            piece = copy.deepcopy(child)
            if (cut_lo, cut_hi) != (start, end):                # only a w:t can straddle
                piece.text = (child.text or '')[cut_lo - start:cut_hi - start]
                piece.set(qn('xml:space'), 'preserve')
            out.append(piece)
        return out

    def _split_and_emit(self, span: MarkerSpan, replacement: str,
                        superscript: Optional[bool]):
        """Replace the runs of *span* with before | citation | after runs.

        The touched runs are split at the child level. The before-run carries
        the first touched run's ``w:rPr`` and every content child of it up to
        the marker; the after-run carries the last touched run's ``w:rPr`` and
        every content child of it from the marker on. A ``w:t`` is cut at the
        marker; a ``w:tab``, ``w:br``, ``w:sym``, ``w:drawing``... survives as
        the element it is, in place. Only the marker text -- and any textless
        child strictly inside it -- is replaced. Whole runs between the first
        and the last touched run lie entirely inside the marker.

        The citation run holds exactly one ``w:t``. Its vertical alignment
        follows *superscript*: set when True, stripped when False, kept from
        the first touched run when None. A superscript citation must not
        bleed into the after-run, which then loses any superscript; otherwise
        the after-run keeps its own. Returns the citation run element.

        The cached :attr:`fields` index is deliberately kept: the span guard
        in ``_locate_span`` has already rejected any touched run that belongs
        to a field, so no field run is removed here, and the runs added carry
        no field characters or code. Every run the identity-keyed index knows
        stays in the tree, so it remains exact -- and rebuilding it per call
        made a 300-token renumber (two calls per token) 20x slower.
        """
        runs, first, first_off, last, last_off = span
        first_r, last_r = runs[first], runs[last]
        before = self._slice_children(first_r, 0, first_off)
        after = self._slice_children(last_r, last_off, None)
        new_runs = []
        if before:
            before_r = self._run_with_rpr(first_r, None)
            before_r.extend(before)
            new_runs.append(before_r)
        cite = self._run_with_rpr(first_r, superscript)
        cite.append(self._text_elem(replacement))
        new_runs.append(cite)
        if after:
            after_r = self._run_with_rpr(last_r, False if superscript else None)
            after_r.extend(after)
            new_runs.append(after_r)
        # Insert after the last touched run (inside its hyperlink, if any),
        # then drop the touched runs from wherever each of them lives.
        for r in reversed(new_runs):
            last_r.addnext(r)
        for r in runs[first:last + 1]:
            r.getparent().remove(r)
        # No invalidate_fields(): see the docstring -- no field run was
        # touched, so the index is still exact and a rebuild per marker
        # would dominate renumbering and export.
        return cite

    def insert_citation_field(self, paragraph, marker_text: str, code: str,
                              result_text: str, superscript: bool, start: int = 0):
        """Replace the first *marker_text* (at or after *start*) with an
        AIREFS citation field.

        Same run surgery as :meth:`replace_marker_by_regex`, but the marker
        becomes the five runs of a complex field whose cached result shows
        *result_text* (formatted like the marker's run, plus ``noProof`` and
        the requested vertical alignment). Returns the result run, or None
        when the marker is not in the paragraph. Raises
        :class:`FieldBoundaryError` if the marker overlaps a field.
        """
        span = self._locate_span(paragraph, marker_text, start)
        if span is None:
            return None
        runs, first, first_off, last, last_off = span
        first_r, last_r = runs[first], runs[last]
        before = self._slice_children(first_r, 0, first_off)
        after = self._slice_children(last_r, last_off, None)
        new_runs = []
        if before:
            before_r = self._run_with_rpr(first_r, None)
            before_r.extend(before)
            new_runs.append(before_r)
        result_run = make_result_run(result_text, first_r, superscript=superscript)
        new_runs.extend(build_field_runs(code, result_run))
        if after:
            after_r = self._run_with_rpr(last_r, False if superscript else None)
            after_r.extend(after)
            new_runs.append(after_r)
        for r in reversed(new_runs):
            last_r.addnext(r)
        for r in runs[first:last + 1]:
            r.getparent().remove(r)
        self.invalidate_fields()          # a new field: the identity index must learn it
        return result_run

    def wrap_run_in_field(self, run_elem, code: str, result_text: str, superscript: bool):
        """Turn one whole run (a superscript citation number list) into an
        AIREFS citation field showing *result_text*. Raises
        :class:`FieldBoundaryError` if the run already belongs to a field."""
        if self.fields.in_field(run_elem):
            raise FieldBoundaryError("run already belongs to a field")
        result_run = make_result_run(result_text, run_elem, superscript=superscript)
        for r in reversed(build_field_runs(code, result_run)):
            run_elem.addnext(r)
        run_elem.getparent().remove(run_elem)
        self.invalidate_fields()
        return result_run

    # ── Bibliography field ───────────────────────────────────────────

    def ensure_bibliography_styles(self):
        """Create the two AI REFs paragraph styles if the document lacks them."""
        styles = self.doc.styles
        names = {s.name for s in styles}
        if BIBLIOGRAPHY_HEADING_STYLE not in names:
            st = styles.add_style(BIBLIOGRAPHY_HEADING_STYLE, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles['Normal']
            st.font.bold = True
            st.font.size = Pt(14)
        if BIBLIOGRAPHY_ENTRY_STYLE not in names:
            st = styles.add_style(BIBLIOGRAPHY_ENTRY_STYLE, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles['Normal']
            st.font.size = Pt(10)

    def _new_entry_paragraphs(self, entries: list[str], bibl_code: str) -> list:
        """Entry paragraphs carrying one AIREFS.BIBL field: begin, code and
        separate in the first, end in the last (all in one when there is a
        single entry). Appended at the end of the body; callers may move them.
        """
        from .docx_fields import _fldchar_run, _instr_run
        paragraphs = []
        n = len(entries)
        for i, entry in enumerate(entries):
            p = self.doc.add_paragraph(style=BIBLIOGRAPHY_ENTRY_STYLE)
            if i == 0:
                p._p.append(_fldchar_run('begin', True))
                p._p.append(_instr_run(bibl_code))
                p._p.append(_fldchar_run('separate'))
            p._p.append(make_result_run(entry, no_proof=False))
            if i == n - 1:
                p._p.append(_fldchar_run('end'))
            paragraphs.append(p)
        return paragraphs

    def write_bibliography_field(self, entries: list[str], bibl_code: str,
                                 heading_text: str = "References"):
        """Append the heading paragraph and the bibliography field."""
        if not entries:
            return
        self.ensure_bibliography_styles()
        heading = self.doc.add_paragraph(heading_text, style=BIBLIOGRAPHY_HEADING_STYLE)
        new_paragraphs = [heading] + self._new_entry_paragraphs(entries, bibl_code)
        self._place_new_paragraphs(new_paragraphs)
        self.invalidate_fields()

    def replace_bibliography_field(self, field, entries: list[str], bibl_code: str) -> int:
        """Replace the paragraphs of an existing AIREFS.BIBL *field* in place.

        The heading and anything after the field are untouched. Returns the
        body index of the first new entry paragraph.
        """
        self.ensure_bibliography_styles()
        old_paragraphs = list(field.paragraphs)
        anchor = old_paragraphs[0]
        new_paragraphs = self._new_entry_paragraphs(entries, bibl_code)
        for p in new_paragraphs:
            anchor.addprevious(p._p)
        for old in old_paragraphs:
            old.getparent().remove(old)
        self.invalidate_fields()
        first = new_paragraphs[0]._p if new_paragraphs else None
        for i, p in enumerate(self.doc.paragraphs):
            if p._p is first:
                return i
        return -1

    def insert_bibliography_after(self, heading_paragraph, entries: list[str], bibl_code: str):
        """Write the bibliography field right after an existing heading paragraph."""
        self.ensure_bibliography_styles()
        self._insert_anchor = heading_paragraph._p.getnext()
        if self._insert_anchor is not None and self._insert_anchor.tag == f'{{{W_NS}}}sectPr':
            self._insert_anchor = None
        self._place_new_paragraphs(self._new_entry_paragraphs(entries, bibl_code))
        self.invalidate_fields()

    def locate_bibliography_heading(self, field, heading_text: str) -> int:
        """Body index of the heading paragraph belonging to a bibliography
        *field*, or -1.

        Rule: the paragraph immediately before the field, if it holds no
        field, is non-empty, shorter than 80 characters and not entry-shaped;
        else the nearest preceding paragraph styled as our heading; else the
        nearest preceding paragraph whose text equals *heading_text*.
        """
        from ..pipeline.existing_citation_parser import BIB_ENTRY_PATTERN
        paragraphs = self.doc.paragraphs
        first_p = field.paragraphs[0] if field.paragraphs else None
        idx = next((i for i, p in enumerate(paragraphs) if p._p is first_p), -1)
        if idx < 0:
            return -1
        cand = idx - 1
        if cand >= 0:
            p = paragraphs[cand]
            text = p.text.strip()
            if (not p._p.findall(f'.//{{{W_NS}}}fldChar') and 0 < len(text) < MAX_HEADING_LENGTH
                    and not BIB_ENTRY_PATTERN.match(text)):
                return cand
        for j in range(idx - 1, -1, -1):
            if paragraphs[j].style is not None and paragraphs[j].style.name == BIBLIOGRAPHY_HEADING_STYLE:
                return j
        wanted = (heading_text or "").strip()
        if wanted:
            for j in range(idx - 1, -1, -1):
                if paragraphs[j].text.strip() == wanted:
                    return j
        return -1

    def validate_before_save(self, expected_cite_fields: Optional[int] = None) -> list[str]:
        """Problems that must abort an export: an incomplete or unparsable
        AIREFS field, field code leaked into visible text, or a citation
        field count that differs from *expected_cite_fields*."""
        from ..pipeline.citation_payload import (
            PayloadError, parse_bibl_code, parse_cite_code,
        )
        self.invalidate_fields()
        problems = []
        idx = self.fields
        for f in idx.fields:
            if not f.kind.startswith('airefs'):
                continue
            if not f.complete:
                problems.append(f"incomplete {f.kind} field")
                continue
            try:
                (parse_cite_code if f.kind == 'airefs_cite' else parse_bibl_code)(f.code)
            except PayloadError as exc:
                problems.append(f"{f.kind} payload does not parse: {exc}")
        for t in self.doc.element.body.iter(f'{{{W_NS}}}t'):
            if t.text and 'ADDIN AIREFS' in t.text:
                problems.append("field code leaked into visible text")
                break
        if expected_cite_fields is not None and idx.airefs_cite != expected_cite_fields:
            problems.append(f"expected {expected_cite_fields} citation fields, found {idx.airefs_cite}")
        return problems

    # ── Insert-mode helpers ─────────────────────────────────────────

    def find_superscript_citation_runs(self, skip_fields: bool = True) -> list[dict]:
        """Find all superscript runs whose text is a citation number list.

        ``numbers`` is the run's text expanded, so a range ``3-5`` yields
        ``[3, 4, 5]``. Expansion is lenient: Word splits a superscript run at
        revision boundaries, so a piece such as ``3-`` (of ``3-5``) is
        reported as ``[3]`` rather than dropped. With ``skip_fields`` (the
        default) a run that belongs to a Word field -- a cached result such
        as an AIREFS.CITE number -- is omitted: field results are rewritten
        through the field API.

        Returns list of dicts with keys:
            paragraph, para_index, run, run_index, numbers (list[int]),
            char_start (int), char_end (int).
        """
        # Local import: the pipeline package depends on this module.
        from ..pipeline.citation_numbers import expand_bracket_numbers
        results = []
        for para_idx, para in enumerate(self.doc.paragraphs):
            char_offset = 0
            for run_idx, run in enumerate(para.runs):
                if run.font.superscript and not (
                        skip_fields and self.fields.in_field(run)):
                    nums = expand_bracket_numbers(run.text, lenient=True)
                    if nums:
                        results.append({
                            'paragraph': para,
                            'para_index': para_idx,
                            'run': run,
                            'run_index': run_idx,
                            'numbers': nums,
                            'char_start': char_offset,
                            'char_end': char_offset + len(run.text),
                        })
                char_offset += len(run.text)
        return results

    def remove_references_section(self, start_para_idx: int,
                                  end_para_idx: Optional[int] = None) -> int:
        """Remove the References section: the heading paragraph and its entries.

        Without *end_para_idx* the section ends where
        ``existing_citation_parser.bibliography_bounds`` says it does (the
        same rule the parser used to read the entries), so an appendix or
        acknowledgements after the list survive. With *end_para_idx* exactly
        ``[start_para_idx, end_para_idx]`` is removed. Refuses to cut through
        a multi-paragraph field. Remembers what followed the section so
        :meth:`append_bibliography` / :meth:`write_bibliography_field` put the
        rebuilt list back in the same place. Returns the number removed.
        """
        from ..pipeline.existing_citation_parser import bibliography_bounds
        paragraphs = self.doc.paragraphs
        if end_para_idx is None:
            end_para_idx = max(start_para_idx, bibliography_bounds(paragraphs, start_para_idx)[1])
        if not (0 <= start_para_idx <= end_para_idx < len(paragraphs)):
            raise ValueError(
                f"invalid References range [{start_para_idx}, {end_para_idx}] "
                f"for a document with {len(paragraphs)} paragraphs")
        self._check_field_boundary(paragraphs, start_para_idx - 1, start_para_idx)
        self._check_field_boundary(paragraphs, end_para_idx, end_para_idx + 1)
        following = paragraphs[end_para_idx]._element.getnext()
        if following is not None and following.tag == f'{{{W_NS}}}sectPr':
            following = None
        body_elem = self.doc.element.body
        for i in range(end_para_idx, start_para_idx - 1, -1):
            body_elem.remove(paragraphs[i]._element)
        self._insert_anchor = following
        removed = end_para_idx - start_para_idx + 1
        logger.info(f"Removed {removed} paragraphs "
                    f"(References section from para {start_para_idx} to {end_para_idx})")
        self.invalidate_fields()
        return removed

    def _check_field_boundary(self, paragraphs, inside_idx: int, outside_idx: int):
        """Raise if one field spans both paragraphs (one to be removed, one kept)."""
        if not (0 <= inside_idx < len(paragraphs) and 0 <= outside_idx < len(paragraphs)):
            return
        inside, outside = paragraphs[inside_idx]._p, paragraphs[outside_idx]._p
        for f in self.fields.fields_in_paragraph(inside):
            if any(q is outside for q in f.paragraphs):
                raise FieldBoundaryError(
                    "the References section boundary cuts through a field; "
                    "refusing to remove part of it")

    def _place_new_paragraphs(self, paragraphs: list):
        """Move freshly appended paragraphs before the remembered anchor, if any."""
        anchor = self._insert_anchor
        if anchor is None:
            return
        for p in paragraphs:
            anchor.addprevious(p._p)
        self._insert_anchor = None

    def append_bibliography(self, entries: list[str]):
        """Write a plain-text bibliography section (legacy, no field).

        Goes where the removed References section was when one was just
        removed, else at the end of the document.
        """
        bib_heading = self.doc.add_paragraph()
        run = bib_heading.add_run("References")
        run.bold = True
        run.font.size = Pt(14)
        bib_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        new_paragraphs = [bib_heading]
        for entry in entries:
            p = self.doc.add_paragraph()
            run = p.add_run(entry)
            run.font.size = Pt(10)
            new_paragraphs.append(p)
        self._place_new_paragraphs(new_paragraphs)
        self.invalidate_fields()

    def save(self, output_path: str):
        """Save the document atomically.

        Writes ``<output_path>.tmp`` beside the target and renames it into
        place, so a failed save (disk full, invalid XML) never leaves a
        half-written or truncated document at *output_path* -- which may be
        the input file itself.
        """
        tmp = f"{output_path}.tmp"
        try:
            self.doc.save(tmp)
            os.replace(tmp, output_path)
        except BaseException:
            try:
                if os.path.exists(tmp):
                    os.remove(tmp)
            except OSError:
                logger.warning(f"Could not remove temporary file {tmp}")
            raise
        logger.info(f"Document saved to {output_path}")
