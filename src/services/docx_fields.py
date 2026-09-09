"""Read Word complex fields (fldChar begin / instrText / separate / result / end).

python-docx has no field API. This module is the only place that understands
field structure; scanners ask a FieldIndex whether a run belongs to a field.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from typing import NamedTuple, Optional

logger = logging.getLogger(__name__)

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
MC_NS = 'http://schemas.openxmlformats.org/markup-compatibility/2006'


def _w(tag: str) -> str:
    """Clark-notation name for a ``w:`` element or attribute."""
    return f'{{{W_NS}}}{tag}'


_MC_FALLBACK = f'{{{MC_NS}}}Fallback'

# Final-view Track Changes semantics. A move made with Track Changes on leaves
# the old copy under w:moveFrom (gone, like w:del) and the new copy under
# w:moveTo (present, like w:ins); both are pending changes.
_REMOVED_WRAPPERS = frozenset({_w('del'), _w('moveFrom')})
_ADDED_WRAPPERS = frozenset({_w('ins'), _w('moveTo')})

CITE_PREFIX = 'ADDIN AIREFS.CITE'
BIBL_PREFIX = 'ADDIN AIREFS.BIBL'
_FOREIGN_MARKERS = ('EN.CITE', 'EN.REFLIST', 'CSL_CITATION', 'CSL_BIBLIOGRAPHY',
                    'ZOTERO_', 'Mendeley', 'CITATION', 'BIBLIOGRAPHY')


def classify_code(code: str) -> str:
    """Classify a field code: ``airefs_cite`` / ``airefs_bibl`` (ours),
    ``foreign`` (another citation manager or Word's native citations) or
    ``other`` (page numbers, TOC, hyperlinks...)."""
    c = (code or '').strip()
    if c.startswith(CITE_PREFIX):
        return 'airefs_cite'
    if c.startswith(BIBL_PREFIX):
        return 'airefs_bibl'
    if c.startswith('ADDIN') and any(m in c for m in _FOREIGN_MARKERS):
        return 'foreign'
    if c.startswith('CITATION') or c.startswith('BIBLIOGRAPHY'):
        return 'foreign'   # Word's native citation fields
    return 'other'


@dataclass
class ComplexField:
    begin: object                                   # w:r holding fldChar begin (or the w:fldSimple)
    code_runs: list = field(default_factory=list)
    separate: Optional[object] = None
    result_runs: list = field(default_factory=list)
    end: Optional[object] = None
    paragraphs: list = field(default_factory=list)  # w:p elements spanned, in order
    depth: int = 0
    deleted: bool = False           # begin run under w:del / w:moveFrom (gone in final view)
    inserted: bool = False          # begin run under w:ins / w:moveTo (present in final view)
    tracked_change: bool = False    # any run of the field under a tracked-change wrapper
    complete: bool = False
    simple: bool = False
    code: str = ''
    kind: str = 'other'
    in_table: bool = False          # begin run inside a w:tbl
    in_text_box: bool = False       # begin run inside a w:txbxContent (DrawingML or VML box)

    @property
    def out_of_flow(self) -> bool:
        """True for table and text-box fields. Their paragraph is not in
        ``doc.paragraphs``, so body-paragraph indexing and first-appearance
        numbering must skip them; the design treats both alike."""
        return self.in_table or self.in_text_box

    @property
    def result_text(self) -> str:
        """Visible result text; paragraphs joined with newlines."""
        from .docx_io import run_text
        chunks, current_p, buf = [], None, []
        for r in self.result_runs:
            p = _enclosing_paragraph(r)
            if current_p is not None and p is not current_p:
                chunks.append(''.join(buf))
                buf = []
            current_p = p
            buf.append(run_text(r))
        chunks.append(''.join(buf))
        return '\n'.join(chunks)

    @property
    def all_runs(self) -> list:
        runs = [self.begin] + list(self.code_runs)
        if self.separate is not None:
            runs.append(self.separate)
        runs += list(self.result_runs)
        if self.end is not None:
            runs.append(self.end)
        return runs


def _enclosing_paragraph(elem):
    p = elem
    while p is not None and p.tag != _w('p'):
        p = p.getparent()
    return p


class RunAncestry(NamedTuple):
    """What the ancestors of a run say about it (see ``run_ancestry``)."""
    skip: bool          # under mc:Fallback: a duplicate of the mc:Choice content
    deleted: bool       # under w:del / w:moveFrom: gone in the final view
    inserted: bool      # under w:ins / w:moveTo: present in the final view
    in_table: bool      # under w:tbl
    in_text_box: bool   # under w:txbxContent

    @property
    def tracked_change(self) -> bool:
        return self.deleted or self.inserted


def run_ancestry(r, stop) -> RunAncestry:
    """Inspect the ancestors of run *r* up to (not including) *stop*.

    The single definition of final-view Track Changes semantics: both the
    field reader here and ``docx_io.iter_all_runs`` derive their ``deleted``
    / ``inserted`` flags from it, so they can never disagree on a run.
    """
    deleted = inserted = in_table = in_text_box = False
    anc = r.getparent()
    while anc is not None and anc is not stop:
        tag = anc.tag
        if tag == _MC_FALLBACK:
            return RunAncestry(True, deleted, inserted, in_table, in_text_box)
        if tag in _REMOVED_WRAPPERS:
            deleted = True
        elif tag in _ADDED_WRAPPERS:
            inserted = True
        elif tag == _w('tbl'):
            in_table = True
        elif tag == _w('txbxContent'):
            in_text_box = True
        anc = anc.getparent()
    return RunAncestry(False, deleted, inserted, in_table, in_text_box)


def _instr_text(run) -> str:
    """Concatenated ``w:instrText`` / ``w:delInstrText`` of a run ('' if none)."""
    return ''.join(child.text or '' for child in run
                   if child.tag in (_w('instrText'), _w('delInstrText')))


def iter_complex_fields(body_elem) -> list[ComplexField]:
    """All fields under *body_elem* in document order (stack-based).

    Reaches body paragraphs, table cells and text boxes (``mc:Fallback``
    duplicates skipped); table and text-box fields are flagged ``in_table`` /
    ``in_text_box`` (together: ``out_of_flow``). Tracked changes follow the
    final view: a field under ``w:del`` or ``w:moveFrom`` is ``deleted``, one
    under ``w:ins`` or ``w:moveTo`` is ``inserted``. ``w:fldSimple`` elements
    are yielded as complete fields whose child runs are the result.
    """
    fields: list[ComplexField] = []
    stack: list[ComplexField] = []
    for node in body_elem.iter(_w('r'), _w('fldSimple')):
        a = run_ancestry(node, body_elem)
        if a.skip:
            continue
        if node.tag == _w('fldSimple'):
            f = ComplexField(begin=node, simple=True, complete=True,
                             code=(node.get(_w('instr')) or '').strip(),
                             depth=len(stack), deleted=a.deleted, inserted=a.inserted,
                             tracked_change=a.tracked_change,
                             in_table=a.in_table, in_text_box=a.in_text_box)
            f.result_runs = list(node.iter(_w('r')))
            f.paragraphs = [_enclosing_paragraph(node)]
            f.kind = classify_code(f.code)
            fields.append(f)
            continue
        fc = node.find(_w('fldChar'))
        p = _enclosing_paragraph(node)
        if fc is not None:
            ftype = fc.get(_w('fldCharType'))
            if ftype == 'begin':
                f = ComplexField(begin=node, depth=len(stack), deleted=a.deleted,
                                 inserted=a.inserted, tracked_change=a.tracked_change,
                                 in_table=a.in_table, in_text_box=a.in_text_box,
                                 paragraphs=[p])
                stack.append(f)
                fields.append(f)
            elif ftype == 'separate' and stack:
                stack[-1].separate = node
                _note_run(stack[-1], p, a)
            elif ftype == 'end' and stack:
                f = stack.pop()
                f.end = node
                f.complete = True
                _note_run(f, p, a)
                f.code = f.code.strip()
                f.kind = classify_code(f.code)
            continue
        if not stack:
            continue
        f = stack[-1]
        _note_run(f, p, a)
        if f.separate is None:          # between begin and separate: field code
            f.code_runs.append(node)
            f.code += _instr_text(node)
        else:                           # between separate and end: cached result
            f.result_runs.append(node)
    for f in stack:                       # unmatched begins
        f.code = f.code.strip()
        f.kind = classify_code(f.code)
        logger.warning("Field begin without a matching end (kind=%s)", f.kind)
    return fields


def _note_run(f: ComplexField, p, a: RunAncestry):
    """Record that a run in paragraph *p* with ancestry *a* belongs to *f*."""
    if p is not None and (not f.paragraphs or f.paragraphs[-1] is not p):
        f.paragraphs.append(p)
    if a.tracked_change:
        f.tracked_change = True


class FieldIndex:
    """Per-document lookup: which runs belong to which field.

    Membership is keyed on lxml element identity, which holds within one
    ``Document`` instance as long as the index (which references every field
    run) is alive. Rebuild the index after any structural edit.
    """

    def __init__(self, doc):
        self.doc = doc
        self.fields = iter_complex_fields(doc.element.body)
        self._role: dict[int, str] = {}
        self._field_of: dict[int, ComplexField] = {}
        # Membership is keyed on element identity: keep every field run alive
        # for as long as the index lives, so a removed run's id can never be
        # recycled by a new run and mistaken for a field run.
        self._pinned = [r for f in self.fields for r in f.all_runs if r is not None]
        for f in self.fields:
            for r in f.all_runs:
                self._field_of.setdefault(id(r), f)
            self._role.setdefault(id(f.begin), 'marker')
            for r in f.code_runs:
                self._role.setdefault(id(r), 'code')
            if f.separate is not None:
                self._role.setdefault(id(f.separate), 'marker')
            for r in f.result_runs:
                self._role.setdefault(id(r), 'result')
            if f.end is not None:
                self._role.setdefault(id(f.end), 'marker')

    # ── counters ────────────────────────────────────────────────
    @property
    def airefs_cite(self) -> int:
        return sum(1 for f in self.fields if f.kind == 'airefs_cite')

    @property
    def airefs_bibl(self) -> int:
        return sum(1 for f in self.fields if f.kind == 'airefs_bibl')

    @property
    def foreign(self) -> int:
        return sum(1 for f in self.fields if f.kind == 'foreign')

    @property
    def out_of_flow(self) -> int:
        """Our own fields inside a table or a text box (``ComplexField.out_of_flow``)."""
        return sum(1 for f in self.fields if f.out_of_flow and f.kind.startswith('airefs'))

    @property
    def in_tables(self) -> int:
        """The plan's name for :attr:`out_of_flow`. Text-box fields are counted
        too: the design treats citations in tables and text boxes alike, and
        both are invisible to ``doc.paragraphs``."""
        return self.out_of_flow

    @property
    def pending_tracked_changes(self) -> bool:
        """Any field with a run under a tracked insertion, deletion or move."""
        return any(f.tracked_change for f in self.fields)

    # ── membership ──────────────────────────────────────────────
    @staticmethod
    def _elem(run):
        return getattr(run, '_r', run)   # accept python-docx Run or raw element

    def in_field(self, run) -> bool:
        return id(self._elem(run)) in self._field_of

    def role(self, run) -> Optional[str]:
        return self._role.get(id(self._elem(run)))

    def field_of(self, run) -> Optional[ComplexField]:
        return self._field_of.get(id(self._elem(run)))

    def fields_in_paragraph(self, paragraph) -> list[ComplexField]:
        p_elem = getattr(paragraph, '_p', paragraph)   # accept Paragraph or raw w:p
        return [f for f in self.fields if any(p is p_elem for p in f.paragraphs)]

    def result_spans(self, paragraph) -> list[tuple[int, int]]:
        """Char spans in paragraph.text covered by top-level field results."""
        from .docx_io import iter_text_runs, run_text
        spans, offset = [], 0
        current = None      # [field, start, end] of the span being extended
        for r in iter_text_runs(paragraph):
            n = len(run_text(r))
            f = self._field_of.get(id(r))
            if f is not None and f.depth == 0 and self._role.get(id(r)) == 'result':
                if current is not None and current[0] is f:
                    current[2] = offset + n
                else:
                    if current is not None:
                        spans.append((current[1], current[2]))
                    current = [f, offset, offset + n]
            offset += n
        if current is not None:
            spans.append((current[1], current[2]))
        return spans


# ── writing fields ──────────────────────────────────────────────────

def make_result_run(text: str, template_r=None, *, superscript: bool = False,
                    no_proof: bool = True):
    """A ``w:r`` for a field's cached result: *template_r*'s ``w:rPr`` (only),
    ``w:noProof`` so Word's checker leaves citation numbers alone, vertical
    alignment per *superscript*, and exactly one ``w:t``."""
    import copy
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = OxmlElement('w:r')
    if template_r is not None:
        src = template_r.find(qn('w:rPr'))
        if src is not None:
            r.append(copy.deepcopy(src))
    rpr = r.get_or_add_rPr()
    for va in rpr.findall(qn('w:vertAlign')):
        rpr.remove(va)
    if no_proof:
        rpr.get_or_add_noProof()
    if superscript:
        rpr.get_or_add_vertAlign().set(qn('w:val'), 'superscript')
    r.append(_text_elem(text))
    return r


def _text_elem(text: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    return t


def _fldchar_run(kind: str, lock: bool = False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = OxmlElement('w:r')
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), kind)
    if lock:
        fc.set(qn('w:fldLock'), '1')
    r.append(fc)
    return r


def _instr_run(code: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = OxmlElement('w:r')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = code
    r.append(it)
    return r


def build_field_runs(code: str, result_run, *, fld_lock: bool = True) -> list:
    """The five runs of a complex field: begin (locked, so F9 never touches
    the result), instruction, separate, *result_run*, end."""
    return [_fldchar_run('begin', fld_lock), _instr_run(code), _fldchar_run('separate'),
            result_run, _fldchar_run('end')]


def rewrite_result(field: ComplexField, text: str, *, superscript: Optional[bool] = None):
    """Set a single-paragraph field's cached result to *text*.

    Word may have re-split the result into several runs (revision
    boundaries); they collapse into the first one, which keeps its ``w:rPr``
    and gets exactly one ``w:t``. *superscript* True/False sets or strips the
    vertical alignment; None keeps it. Not for multi-paragraph fields (the
    bibliography is replaced paragraph-wise instead).
    """
    from docx.oxml.ns import qn
    if not field.result_runs:
        keep = make_result_run(text, superscript=bool(superscript))
        anchor = field.separate if field.separate is not None else field.begin
        anchor.addnext(keep)
        field.result_runs = [keep]
        return keep
    keep = field.result_runs[0]
    for r in field.result_runs[1:]:
        r.getparent().remove(r)
    for child in list(keep):
        if child.tag != qn('w:rPr'):
            keep.remove(child)
    keep.append(_text_elem(text))
    if superscript is not None:
        rpr = keep.get_or_add_rPr()
        for va in rpr.findall(qn('w:vertAlign')):
            rpr.remove(va)
        if superscript:
            rpr.get_or_add_vertAlign().set(qn('w:val'), 'superscript')
    field.result_runs = [keep]
    return keep


def unwrap_field(field: ComplexField):
    """Remove a field's markers and code, keeping its visible result runs as
    plain text (the inverse of wrapping). Used when the user typed a marker
    into an unresolved [?] field: the marker is then processed like any other."""
    from docx.oxml.ns import qn
    for r in [field.begin, *field.code_runs, field.separate, field.end]:
        if r is not None and r.getparent() is not None:
            r.getparent().remove(r)
    for r in field.result_runs:
        rpr = r.find(qn('w:rPr'))
        if rpr is not None:
            for np_ in rpr.findall(qn('w:noProof')):
                rpr.remove(np_)
    field.complete = False


def rewrite_code(field: ComplexField, code: str):
    """Replace the field's instruction text with *code* (one code run)."""
    from docx.oxml.ns import qn
    if not field.code_runs:
        run = _instr_run(code)
        field.begin.addnext(run)
        field.code_runs = [run]
    else:
        keep = field.code_runs[0]
        for r in field.code_runs[1:]:
            r.getparent().remove(r)
        for child in list(keep):
            if child.tag in (qn('w:instrText'), qn('w:delInstrText')):
                keep.remove(child)
        it = _instr_run(code).find(qn('w:instrText'))
        keep.append(it)
        field.code_runs = [keep]
    field.code = code.strip()
    field.kind = classify_code(field.code)
