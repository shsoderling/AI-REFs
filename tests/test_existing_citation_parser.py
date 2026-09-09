"""Tests for parsing pre-existing citations from a document."""

import pytest
from docx import Document

from src.pipeline.existing_citation_parser import ExistingCitationParser
from src.services.docx_io import DocxHandler


def make_doc(tmp_path, body_paras, bib_paras, heading="References",
             name="doc.docx"):
    doc = Document()
    for item in body_paras:
        if isinstance(item, str):
            doc.add_paragraph(item)
        else:
            # list of (text, is_superscript) runs
            para = doc.add_paragraph()
            for text, sup in item:
                run = para.add_run(text)
                run.font.superscript = sup
    doc.add_paragraph(heading)
    for text in bib_paras:
        doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def analyze(path):
    return ExistingCitationParser(DocxHandler(path)).analyze()


class TestHeadingDetection:
    @pytest.mark.parametrize("heading", [
        "References", "REFERENCES", "Bibliography", "Literature Cited",
        "Works Cited",
    ])
    def test_heading_variants(self, tmp_path, heading):
        path = make_doc(tmp_path, ["Body."], ["1. Smith J. Title. J. 2020."],
                        heading=heading)
        result = analyze(path)
        assert result.references_heading_para_idx == 1
        assert len(result.bib_entries) == 1


class TestBibliographyEntryFormats:
    @pytest.mark.parametrize("entry", [
        "1. Smith J. A title. Journal. 2020.",
        "1) Smith J. A title. Journal. 2020.",
        "[1] Smith J. A title. Journal. 2020.",
    ])
    def test_entry_number_formats(self, tmp_path, entry):
        path = make_doc(tmp_path, ["Body [1]."], [entry])
        result = analyze(path)
        assert 1 in result.bib_entries
        assert result.bib_entries[1].body.startswith("Smith J.")

    def test_doi_pmid_year_extraction(self, tmp_path):
        path = make_doc(tmp_path, ["Body [1]."], [
            "1. Smith J. A title. Journal. 2021;12:1-10. "
            "doi:10.1234/abc.5678 PMID: 31234567",
        ])
        entry = analyze(path).bib_entries[1]
        assert entry.doi == "10.1234/abc.5678"
        assert entry.pmid == "31234567"
        assert entry.year == 2021

    def test_body_stored_for_renumbering(self, tmp_path):
        path = make_doc(tmp_path, ["Body [1]."],
                        ["1. Smith J. A title. Journal. 2020."])
        entry = analyze(path).bib_entries[1]
        assert entry.body == "Smith J. A title. Journal. 2020."


class TestInTextCitations:
    def _bib(self, n):
        return [f"{i}. Author{i} A. Title. J. 2020." for i in range(1, n + 1)]

    def test_bracket_citations_with_semicolon(self, tmp_path):
        """Regression for pattern drift: [1; 3] must be parsed."""
        path = make_doc(tmp_path, ["As shown [1; 3]."], self._bib(3))
        cites = analyze(path).in_text_citations[0]
        assert sorted(c.number for c in cites) == [1, 3]

    def test_bracket_range_expansion(self, tmp_path):
        path = make_doc(tmp_path, ["As shown [1-3]."], self._bib(3))
        cites = analyze(path).in_text_citations[0]
        assert sorted(c.number for c in cites) == [1, 2, 3]

    def test_phantom_bracket_number_excluded(self, tmp_path):
        """[2020] (a year) has no bib entry — it must not become a citation."""
        path = make_doc(tmp_path, ["The big study [2020] and real cite [1]."],
                        self._bib(2))
        cites = analyze(path).in_text_citations[0]
        assert sorted(c.number for c in cites) == [1]

    def test_superscript_citation_detected(self, tmp_path):
        path = make_doc(
            tmp_path,
            [[("As shown", False), ("1,2", True), (".", False)]],
            self._bib(2),
        )
        cites = analyze(path).in_text_citations[0]
        assert sorted(c.number for c in cites) == [1, 2]
        assert all(c.is_superscript for c in cites)

    def test_non_citation_superscript_excluded(self, tmp_path):
        """Regression: the '2+' in Ca2+ must not be recorded as citation 2."""
        path = make_doc(
            tmp_path,
            [[("Ca", False), ("2+", True), (" signaling", False)]],
            self._bib(3),
        )
        assert 0 not in analyze(path).in_text_citations

    def test_phantom_superscript_number_excluded(self, tmp_path):
        """A superscript number with no bib entry must be ignored."""
        path = make_doc(
            tmp_path,
            [[("Footnote", False), ("9", True)]],
            self._bib(2),
        )
        assert 0 not in analyze(path).in_text_citations

    def test_numbers_inside_markers_skipped(self, tmp_path):
        path = make_doc(tmp_path, ["New claim (REF) but old cite [1]."],
                        self._bib(1))
        cites = analyze(path).in_text_citations[0]
        assert [c.number for c in cites] == [1]


class TestStyleDetection:
    def test_bracket_style_detected(self, tmp_path):
        path = make_doc(tmp_path, ["A [1].", "B [2]."],
                        ["1. A. T. J. 2020.", "2. B. T. J. 2021."])
        assert analyze(path).detected_style_is_superscript is False

    def test_superscript_style_detected(self, tmp_path):
        path = make_doc(
            tmp_path,
            [[("A", False), ("1", True)], [("B", False), ("2", True)]],
            ["1. A. T. J. 2020.", "2. B. T. J. 2021."],
        )
        assert analyze(path).detected_style_is_superscript is True


class TestNoExistingCitations:
    def test_fresh_document(self, tmp_path):
        doc = Document()
        doc.add_paragraph("Just a fresh document (REF).")
        path = tmp_path / "fresh.docx"
        doc.save(str(path))
        result = analyze(str(path))
        assert not result.has_existing_citations
        assert result.references_heading_para_idx == -1
