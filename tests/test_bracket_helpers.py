"""Tests for bracket-citation helpers and renumbering application on real docs."""

import pytest
from docx import Document

from src.models.existing_refs import ExistingCitationMap
from src.pipeline.renumber_apply import (
    apply_renumbering, expand_bracket_numbers, format_bracket_numbers,
    renumber_bracket_group,
)
from src.services.docx_io import DocxHandler


class TestExpandBracketNumbers:
    @pytest.mark.parametrize("text,expected", [
        ("1", [1]),
        ("1,2,3", [1, 2, 3]),
        ("1, 2, 3", [1, 2, 3]),
        ("1-3", [1, 2, 3]),
        ("1–3", [1, 2, 3]),          # en-dash
        ("1; 3", [1, 3]),
        ("2, 5-7", [2, 5, 6, 7]),
    ])
    def test_expand(self, text, expected):
        assert expand_bracket_numbers(text) == expected


class TestFormatBracketNumbers:
    @pytest.mark.parametrize("nums,expected", [
        ([1], "1"),
        ([1, 2], "1, 2"),            # two consecutive: no range
        ([1, 2, 3], "1-3"),
        ([2, 3, 5], "2, 3, 5"),
        ([1, 2, 3, 7, 8, 9], "1-3, 7-9"),
        ([], ""),
    ])
    def test_format(self, nums, expected):
        assert format_bracket_numbers(nums) == expected


class TestRenumberBracketGroup:
    def test_simple_map(self):
        assert renumber_bracket_group("1", {1: 2}) == "2"

    def test_range_remap_no_false_range(self):
        """[1-3] mapped to 2,3,5 must NOT become 2-5."""
        assert renumber_bracket_group("1-3", {1: 2, 2: 3, 3: 5}) == "2, 3, 5"

    def test_dedup_after_merge(self):
        assert renumber_bracket_group("1, 2", {1: 4, 2: 4}) == "4"


def _make_doc(tmp_path, paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    return str(path)


def _existing_map(refs_heading):
    return ExistingCitationMap(references_heading_para_idx=refs_heading)


class TestApplyRenumberingBrackets:
    def test_cascade_regression(self, tmp_path):
        """Regression for the cascade bug: {1->2, 2->3} on '[1] ... [2]'.

        Sequential in-place replacement would turn [1] into [2] and then the
        second rule would re-match the freshly inserted [2], producing
        '[3] ... [2]'.  The two-phase sentinel replacement must yield
        '[2] ... [3]'.
        """
        path = _make_doc(tmp_path, [
            "Earlier work [1] showed X, and later work [2] showed Y.",
            "References",
        ])
        handler = DocxHandler(path)
        apply_renumbering(handler, _existing_map(refs_heading=1), {1: 2, 2: 3})
        text = handler.get_paragraphs()[0].text
        assert text == "Earlier work [2] showed X, and later work [3] showed Y."

    def test_swap_does_not_collide(self, tmp_path):
        path = _make_doc(tmp_path, ["See [1] and [2].", "References"])
        handler = DocxHandler(path)
        apply_renumbering(handler, _existing_map(1), {1: 2, 2: 1})
        assert handler.get_paragraphs()[0].text == "See [2] and [1]."

    def test_range_in_paragraph(self, tmp_path):
        path = _make_doc(tmp_path, ["Shown previously [1-3].", "References"])
        handler = DocxHandler(path)
        apply_renumbering(handler, _existing_map(1), {1: 2, 2: 3, 3: 5})
        assert handler.get_paragraphs()[0].text == "Shown previously [2, 3, 5]."

    def test_superscript_bracket_citation_keeps_its_superscript(self, tmp_path):
        """Superscript-bracket styles: '[3]' is not citation-shaped, so it
        takes the bracket path, which rewrites the token in place and must
        leave its vertical alignment alone."""
        doc = Document()
        para = doc.add_paragraph()
        para.add_run("See")
        para.add_run("[3]").font.superscript = True
        para.add_run(".")
        doc.add_paragraph("References")
        path = tmp_path / "supbr.docx"
        doc.save(str(path))
        handler = DocxHandler(str(path))
        apply_renumbering(handler, _existing_map(1), {3: 7})
        para = handler.get_paragraphs()[0]
        assert [(r.text, r.font.superscript) for r in para.runs] == [
            ("See", None), ("[7]", True), (".", None)]

    def test_references_section_untouched(self, tmp_path):
        path = _make_doc(tmp_path, [
            "Body cite [1].",
            "References",
            "1. Old entry [1] mention.",
        ])
        handler = DocxHandler(path)
        apply_renumbering(handler, _existing_map(1), {1: 9})
        assert handler.get_paragraphs()[0].text == "Body cite [9]."
        assert handler.get_paragraphs()[2].text == "1. Old entry [1] mention."


class TestApplyRenumberingSuperscript:
    def _make_superscript_doc(self, tmp_path, runs, trailing_paras=("References",)):
        """runs: list of (text, is_superscript) composing one paragraph."""
        doc = Document()
        para = doc.add_paragraph()
        for text, sup in runs:
            run = para.add_run(text)
            run.font.superscript = sup
        for t in trailing_paras:
            doc.add_paragraph(t)
        path = tmp_path / "sup.docx"
        doc.save(str(path))
        return str(path)

    def test_superscript_citation_renumbered(self, tmp_path):
        path = self._make_superscript_doc(
            tmp_path, [("As shown", False), ("1,2", True), (".", False)])
        handler = DocxHandler(path)
        apply_renumbering(handler, _existing_map(1), {1: 3, 2: 4})
        para = handler.get_paragraphs()[0]
        sup_texts = [r.text for r in para.runs if r.font.superscript]
        assert sup_texts == ["3,4"]

    def test_non_citation_superscript_untouched(self, tmp_path):
        """The '2+' in Ca2+ is superscript but not a citation — leave it alone."""
        path = self._make_superscript_doc(
            tmp_path, [("Ca", False), ("2+", True), (" influx", False),
                       ("2", True)])
        handler = DocxHandler(path)
        apply_renumbering(handler, _existing_map(1), {2: 7})
        para = handler.get_paragraphs()[0]
        sup_texts = [r.text for r in para.runs if r.font.superscript]
        assert sup_texts == ["2+", "7"]

    def test_no_cascade_in_superscript(self, tmp_path):
        path = self._make_superscript_doc(
            tmp_path, [("A", False), ("1", True), (" and B", False), ("2", True)])
        handler = DocxHandler(path)
        apply_renumbering(handler, _existing_map(1), {1: 2, 2: 3})
        para = handler.get_paragraphs()[0]
        sup_texts = [r.text for r in para.runs if r.font.superscript]
        assert sup_texts == ["2", "3"]
