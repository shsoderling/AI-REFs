"""Round-trip tests: the app's own export must re-parse cleanly in insert mode.

The core re-upload workflow is: fresh export -> user adds new (REF) markers ->
re-upload -> insert mode merges and renumbers.  That only works if
ExistingCitationParser fully recovers what the export wrote.
"""

import pytest
from docx import Document

from src.models.citation import Author, CitationCandidate
from src.models.project import CitationStyle
from src.pipeline.bib_format import format_bib_entry
from src.pipeline.existing_citation_parser import ExistingCitationParser
from src.pipeline.renumbering import NewMarkerInfo, compute_renumbering
from src.services.docx_io import DocxHandler


def make_citation(i):
    return CitationCandidate(
        pmid=f"3000000{i}",
        doi=f"10.1234/test.{i}",
        title=f"Important Study Number {i}",
        authors=[Author(last_name=f"Author{i}", initials="A")],
        year=2018 + i,
        journal=f"Journal of Tests {i}",
        journal_abbrev=f"J Test {i}",
        volume=str(10 + i),
        pages=f"{i}00-{i}10",
    )


def export_like_fresh(path, body_paras, citations, style, superscript=False):
    """Mimic what _do_fresh_export writes: cited body + References section."""
    doc = Document()
    for text, cite_nums in body_paras:
        para = doc.add_paragraph()
        para.add_run(text)
        if cite_nums:
            if superscript:
                run = para.add_run(",".join(str(n) for n in cite_nums))
                run.font.superscript = True
            else:
                para.add_run(f"[{', '.join(str(n) for n in cite_nums)}]")
            para.add_run(".")
    doc.save(str(path))

    handler = DocxHandler(str(path))
    entries = [
        format_bib_entry(c, i, style) for i, c in enumerate(citations, 1)
    ]
    handler.append_bibliography(entries)
    handler.save(str(path))
    return str(path)


@pytest.mark.parametrize("superscript", [False, True],
                         ids=["bracket", "superscript"])
def test_fresh_export_reparses_completely(tmp_path, superscript):
    citations = [make_citation(i) for i in range(1, 4)]
    path = export_like_fresh(
        tmp_path / "out.docx",
        [("First finding ", [1]), ("Second finding ", [2, 3])],
        citations,
        CitationStyle.NIH_GRANT,
        superscript=superscript,
    )

    result = ExistingCitationParser(DocxHandler(path)).analyze()

    assert result.has_existing_citations
    assert result.references_heading_para_idx == 2
    assert sorted(result.bib_entries.keys()) == [1, 2, 3]
    # Identifiers recovered for dedup
    for i in range(1, 4):
        assert result.bib_entries[i].pmid == f"3000000{i}"
        assert result.bib_entries[i].doi == f"10.1234/test.{i}"
    # All in-text citations recovered
    found = sorted(
        c.number for cites in result.in_text_citations.values() for c in cites
    )
    assert found == [1, 2, 3]
    assert result.detected_style_is_superscript == superscript


def test_reupload_dedups_new_candidate_against_own_export(tmp_path):
    """A re-found paper (same PMID) must reuse the existing number."""
    citations = [make_citation(1), make_citation(2)]
    path = export_like_fresh(
        tmp_path / "out.docx",
        [("Alpha ", [1]), ("Beta ", [2])],
        citations,
        CitationStyle.NIH_GRANT,
    )
    existing = ExistingCitationParser(DocxHandler(path)).analyze()

    # New marker resolves to paper 2 (found again) and one genuinely new paper
    refound = CitationCandidate(pmid=citations[1].pmid, title="Refound copy")
    brand_new = make_citation(7)
    new = [NewMarkerInfo(para_index=0, char_offset=0,
                         citations=[refound, brand_new])]
    result = compute_renumbering(existing, new)

    assert len(result.assignments) == 3  # not 4 — refound merged with entry 2
    assert result.number_for_candidate(refound) == result.renumber_map[2]


def test_format_bib_entry_always_includes_identifiers():
    """DOI and PMID must be emitted for every numeric style (round-trip keys)."""
    c = make_citation(1)
    for style in (CitationStyle.NIH_GRANT, CitationStyle.NATURE,
                  CitationStyle.SCIENCE):
        entry = format_bib_entry(c, 1, style)
        assert f"doi:{c.doi}" in entry
        assert f"PMID: {c.pmid}" in entry
