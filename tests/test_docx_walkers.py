"""Tests for the run walkers in docx_io: the text-run walker must reproduce
``paragraph.text`` exactly; the all-run walker must reach nested runs and
skip mc:Fallback duplicates."""

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK

from src.services.docx_io import iter_text_runs, iter_all_runs, run_text
from tests.fixture_builders import DocBuilder, make_run

FIXTURES = sorted((Path(__file__).resolve().parents[1] / "fixtures").glob("*.docx"))


def _edge_doc(tmp_path):
    b = DocBuilder()
    p = b.paragraph("Plain ")
    b.add_hyperlink(p, "https://x.org", "link")
    b.add_text(p, " then ")
    b.add_field(p, code=' ADDIN AIREFS.CITE {"v":1} ', result="3", superscript=True)
    b.add_text(p, " ins", wrap="ins")
    b.add_text(p, " del", wrap="del")
    b.add_inline_sdt(p, "TAG", " sdt")
    b.add_multi_wt_run(p, ["M1", "M2"])
    return b.save(tmp_path / "edge.docx")


@pytest.mark.parametrize("path", FIXTURES, ids=[p.name for p in FIXTURES])
def test_text_runs_reproduce_paragraph_text_on_fixtures(path):
    for para in Document(str(path)).paragraphs:
        assert "".join(run_text(r) for r in iter_text_runs(para)) == para.text


def test_text_runs_reproduce_paragraph_text_on_edge_doc(tmp_path):
    para = Document(str(_edge_doc(tmp_path))).paragraphs[0]
    texts = [run_text(r) for r in iter_text_runs(para)]
    assert "".join(texts) == para.text
    assert "link" in texts            # hyperlink run is included
    assert " sdt" not in texts        # sdt is invisible to paragraph.text, so also to us
    assert "M1M2" in texts            # multi-w:t run text is concatenated


def test_all_runs_covers_nested_and_marks_deleted(tmp_path):
    para = Document(str(_edge_doc(tmp_path))).paragraphs[0]
    infos = list(iter_all_runs(para))
    texts = [run_text(i.elem) for i in infos]
    assert " ins" in texts and " sdt" in texts and "link" in texts
    deleted = [i for i in infos if i.deleted]
    assert len(deleted) == 1 and deleted[0].elem.xpath("string(w:delText)") == " del"
    inserted = [i for i in infos if i.inserted]
    assert len(inserted) == 1 and run_text(inserted[0].elem) == " ins"


def test_all_runs_skips_mc_fallback(tmp_path):
    b = DocBuilder()
    p = b.paragraph("Body")

    def fill(inner_p):
        inner_p.append(make_run("boxed"))

    b.add_text_box(p, fill)
    para = Document(str(b.save(tmp_path / "tb.docx"))).paragraphs[0]
    texts = [run_text(i.elem) for i in iter_all_runs(para)]
    assert texts.count("boxed") == 1


def test_run_text_matches_python_docx_for_breaks_and_tabs(tmp_path):
    b = DocBuilder()
    p = b.paragraph()
    r = p.add_run("a")
    r.add_break(WD_BREAK.PAGE)      # page break: no text, unlike a line break
    r.add_tab()
    r.add_break(WD_BREAK.LINE)
    r.add_text("b")
    para = Document(str(b.save(tmp_path / "br.docx"))).paragraphs[0]
    run = para.runs[0]._r
    assert run_text(run) == run.text == "a\t\nb"
    assert "".join(run_text(x) for x in iter_text_runs(para)) == para.text


def test_all_runs_marks_tracked_moves(tmp_path):
    """A move made with Track Changes on leaves the old copy under w:moveFrom
    and the new copy under w:moveTo. Final view: moveFrom is gone (like
    w:del), moveTo is present (like w:ins). Moved-from text keeps w:t, so
    run_text alone cannot tell the stale copy from live text -- only the
    deleted flag can."""
    b = DocBuilder()
    p = b.paragraph("A")
    b.add_text(p, "B", wrap="moveFrom")
    b.add_text(p, "C", wrap="moveTo")
    para = Document(str(b.save(tmp_path / "mv.docx"))).paragraphs[0]
    by_text = {run_text(i.elem): (i.deleted, i.inserted) for i in iter_all_runs(para)}
    assert by_text == {"A": (False, False), "B": (True, False), "C": (False, True)}


def test_all_runs_and_field_index_agree_on_final_view(tmp_path):
    """There is one definition of 'gone in the final view': every run of a
    field, as seen through iter_all_runs, carries the same deleted/inserted
    flags as the field itself in FieldIndex, for all four wrapper kinds."""
    from src.services.docx_fields import FieldIndex
    from tests.fixture_builders import TRACKED_CHANGE_KINDS

    b = DocBuilder()
    paras = []
    for kind in TRACKED_CHANGE_KINDS:
        p = b.paragraph("")
        b.add_field(p, code=f' ADDIN AIREFS.CITE {{"k":"{kind}"}} ', result="1", wrap=kind)
        paras.append(p)
    doc = Document(str(b.save(tmp_path / "agree.docx")))
    idx = FieldIndex(doc)
    assert [f.deleted for f in idx.fields] == [False, True, True, False]
    assert [f.inserted for f in idx.fields] == [True, False, False, True]
    for para, f in zip(doc.paragraphs, idx.fields):
        infos = list(iter_all_runs(para))
        assert infos, "field runs must be reachable through iter_all_runs"
        assert all(idx.field_of(i.elem) is f for i in infos)
        assert {(i.deleted, i.inserted) for i in infos} == {(f.deleted, f.inserted)}
