"""DocxHandler.save writes atomically: temp file beside the target, then rename."""
import os

import pytest
from docx import Document

import src.services.docx_io as docx_io
from src.services.docx_io import DocxHandler


def test_save_is_atomic_and_leaves_no_tmp(tmp_path, monkeypatch):
    calls = []
    real_replace = os.replace
    monkeypatch.setattr(docx_io.os, "replace",
                        lambda a, b: calls.append((str(a), str(b))) or real_replace(a, b))
    src = tmp_path / "in.docx"
    Document().save(str(src))
    h = DocxHandler(str(src))
    h.doc.add_paragraph("added")
    out = tmp_path / "out.docx"
    h.save(str(out))
    assert Document(str(out)).paragraphs[-1].text == "added"
    assert calls and calls[0][0].endswith(".tmp") and calls[0][1] == str(out)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["in.docx", "out.docx"]


def test_save_over_input_path_works(tmp_path):
    src = tmp_path / "in.docx"
    Document().save(str(src))
    h = DocxHandler(str(src))
    h.doc.add_paragraph("again")
    h.save(str(src))
    assert Document(str(src)).paragraphs[-1].text == "again"


def test_failed_save_leaves_target_untouched(tmp_path, monkeypatch):
    src = tmp_path / "in.docx"
    Document().save(str(src))
    out = tmp_path / "out.docx"
    out.write_bytes(b"previous")
    h = DocxHandler(str(src))
    monkeypatch.setattr(h.doc, "save", lambda path: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        h.save(str(out))
    assert out.read_bytes() == b"previous"
    assert not (tmp_path / "out.docx.tmp").exists()
