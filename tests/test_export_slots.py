"""export_slots: per-marker actions, plus an end-to-end DOCX export through docx_export."""

from docx import Document

from src.models.citation import Author, CitationCandidate, PLACEHOLDER_TITLE_PREFIX
from src.models.evidence import EvidenceRecord, ReviewDecision
from src.models.project import ProjectState, CitationStyle
from src.pipeline.docx_export import export_fresh
from src.pipeline.document_parser import DocumentParser
from src.pipeline.export_slots import (
    ACTION_CITE, ACTION_LEAVE, ACTION_UNRESOLVED,
    collect_marker_slots, slot_for_docx_marker, action_for_unmatched,
)
from src.pipeline.marker_locator import MarkerLocator
from src.pipeline.renumber_plan import build_renumber_plan
from src.services.docx_io import DocxHandler


def cand(pmid, title=None, year=2020):
    return CitationCandidate(pmid=pmid, title=title or f"Paper {pmid}", year=year, journal_abbrev="J",
                             journal="Journal", volume="1", pages="1-2",
                             authors=[Author(last_name=f"Author{pmid}", initials="A")])


def build_project(tmp_path, paragraphs, embed=False):
    path = tmp_path / "in.docx"
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))
    project = ProjectState(input_docx_path=str(path))
    project.settings.embed_citation_fields = embed
    handler = DocxHandler(str(path))
    sentences = DocumentParser(handler, marker_config=project.marker_config).parse()
    MarkerLocator(project.marker_config).locate(sentences)
    project.sentences = sentences
    project.run_marker_config = project.marker_config   # what the orchestrator records
    return project


def test_actions_per_marker_kind(tmp_path):
    project = build_project(tmp_path, [
        "A (REF) B (PMID: 1) C (Smith 2020) D (REFS) E (PMC2000).",
    ])
    (s,) = project.sentences
    assert len(s.markers) == 5 and s.slot_count == 5
    ev = EvidenceRecord(sentence_id=s.id)
    # slot 0 REF: cited; slot 1 PMID: cited; slot 2 author-year: nothing found;
    # slot 3 REFS: placeholder only; slot 4 PMC: found but user removed it
    ev.selected = [cand("10"), cand("1"), CitationCandidate(title=f"{PLACEHOLDER_TITLE_PREFIX} x)")]
    ev.slot_sizes = [1, 1, 0, 1, 0]
    ev.review_decision = ReviewDecision.ACCEPTED
    project.evidence_map[s.id] = ev

    slots = collect_marker_slots(project)[s.paragraph_index]
    assert [x.action for x in slots] == [ACTION_CITE, ACTION_CITE, ACTION_LEAVE, ACTION_UNRESOLVED, ACTION_LEAVE]
    assert [c.pmid for c in slots[0].citations] == ["10"]
    assert slots[3].citations == []  # placeholder filtered out

    # "Leave unchanged" -> everything left as written
    ev.review_decision = ReviewDecision.SKIPPED
    slots = collect_marker_slots(project)[s.paragraph_index]
    assert {x.action for x in slots} == {ACTION_LEAVE}
    assert all(x.is_skipped for x in slots)

    # pending (or rejected) sentence -> REF/REFS become [?], suggested stay
    for decision in (ReviewDecision.PENDING, ReviewDecision.REJECTED):
        ev.review_decision = decision
        slots = collect_marker_slots(project)[s.paragraph_index]
        assert [x.action for x in slots] == [ACTION_UNRESOLVED, ACTION_LEAVE, ACTION_LEAVE, ACTION_UNRESOLVED, ACTION_LEAVE]

    assert slot_for_docx_marker({0: slots}, 0, 7) is None
    assert slot_for_docx_marker({0: slots}, 0, 1, "(PMID: 1)") is slots[1]
    assert slot_for_docx_marker({0: slots}, 0, 1, "(PMID: 2)") is None      # text must agree
    assert action_for_unmatched("SUGGESTED") == ACTION_LEAVE and action_for_unmatched("REF") == ACTION_UNRESOLVED


def test_all_refs_sentence_cites_its_whole_list_at_every_marker(tmp_path):
    project = build_project(tmp_path, ["Both (REFS) and (REFS) here."])
    (s,) = project.sentences
    assert s.slot_count == 1 and not s.searched_per_marker
    project.evidence_map[s.id] = EvidenceRecord(
        sentence_id=s.id, selected=[cand("1"), cand("2")], slot_sizes=[2],
        review_decision=ReviewDecision.ACCEPTED)
    slots = collect_marker_slots(project)[0]
    assert [[c.pmid for c in x.citations] for x in slots] == [["1", "2"], ["1", "2"]]


def test_plan_gives_no_number_to_markers_left_as_written(tmp_path):
    project = build_project(tmp_path, ["A (REF) then (Jones 2019) then (REF)."])
    (s,) = project.sentences
    project.evidence_map[s.id] = EvidenceRecord(
        sentence_id=s.id, selected=[cand("1"), cand("2")], slot_sizes=[1, 0, 1],
        review_decision=ReviewDecision.ACCEPTED)
    plan = build_renumber_plan(DocxHandler(project.input_docx_path), project)
    assert [m["text"] for m in plan.markers] == ["(REF)", "(Jones 2019)", "(REF)"]
    assert plan.marker_actions == [ACTION_CITE, ACTION_LEAVE, ACTION_CITE]
    assert plan.written_count == 2
    assert sorted(plan.renumber_result.assignments) == [1, 2]


def test_end_to_end_docx_export(tmp_path):
    """Replace cited markers, write [?] for unresolved (REF), keep unconfirmed suggestions verbatim."""
    project = build_project(tmp_path, [
        "First claim (REF). Second claim (PMID: 1) and third (Smith et al. 2020).",
        "Fourth (REF) fourth-b (REF) fifth (PMC7000000).",
        "Sixth keeps its text (Jones 2019).",
    ])
    project.settings.citation_style = CitationStyle.VANCOUVER
    by_text = {s.raw_text: s for s in project.sentences}

    s1 = by_text["First claim (REF)."]
    project.evidence_map[s1.id] = EvidenceRecord(
        sentence_id=s1.id, selected=[cand("100")], slot_sizes=[1], review_decision=ReviewDecision.ACCEPTED)

    s2 = by_text["Second claim (PMID: 1) and third (Smith et al. 2020)."]
    project.evidence_map[s2.id] = EvidenceRecord(
        sentence_id=s2.id, selected=[cand("1")], slot_sizes=[1, 0], review_decision=ReviewDecision.ACCEPTED)

    s3 = by_text["Fourth (REF) fourth-b (REF) fifth (PMC7000000)."]
    # slot 0 empty, slot 1 cited, slot 2 cited (same paper as sentence 1 -> reuses number 1)
    project.evidence_map[s3.id] = EvidenceRecord(
        sentence_id=s3.id, selected=[cand("200"), cand("100")], slot_sizes=[0, 1, 1],
        review_decision=ReviewDecision.ACCEPTED)

    s4 = by_text["Sixth keeps its text (Jones 2019)."]
    project.evidence_map[s4.id] = EvidenceRecord(sentence_id=s4.id, review_decision=ReviewDecision.SKIPPED)

    out = tmp_path / "out.docx"
    stats = export_fresh(project, str(out))

    paras = [p.text for p in Document(str(out)).paragraphs]
    # Vancouver CSL renders in-text citations as (n)
    assert paras[0] == "First claim (1). Second claim (2) and third (Smith et al. 2020)."
    assert paras[1] == "Fourth [?] fourth-b (3) fifth (1)."
    assert paras[2] == "Sixth keeps its text (Jones 2019)."
    assert "References" in paras
    bib = [p for p in paras if p[:2] in ("1.", "2.", "3.")]
    assert bib[0].startswith("1. Author100")
    assert (stats.total_markers, stats.resolved_markers, stats.unresolved_markers) == (7, 4, 1)
    assert (stats.left_unverified, stats.skipped_markers, stats.unmatched_markers) == (1, 1, 0)
    assert stats.unresolved_sentence_ids == [s3.id]
    assert any("left unverified" in line for line in stats.summary_lines())

    # Same document with tracked fields: markers left as written get no field
    project.settings.embed_citation_fields = True
    out2 = tmp_path / "out2.docx"
    stats2 = export_fresh(project, str(out2))
    assert stats2.fields_written == 5 and stats2.unresolved_fields == 1
    h = DocxHandler(str(out2))
    assert sum(1 for f in h.fields.fields if f.kind == "airefs_cite") == 5
    assert [p.text for p in h.get_paragraphs()][:3] == paras[:3]

    # A project saved by the old app version records no run configuration:
    # export then scans with the (REF)/(REFS)-only grammar even though the
    # current settings have detection on, so author-year text is never touched.
    project.settings.embed_citation_fields = False
    project.run_marker_config = None
    for s in project.sentences:
        s.markers = [m for m in s.markers if m.kind.value in ("REF", "REFS")]
        s.marker_types = [m.kind for m in s.markers]
        s.marker_count = len(s.markers)
        s.marker_type = s.markers[0].kind if s.markers else None
    for ev in project.evidence_map.values():
        ev.slot_sizes = []
    stats3 = export_fresh(project, str(out))
    paras = [p.text for p in Document(str(out)).paragraphs]
    assert paras[0].startswith("First claim (1). Second claim (PMID: 1)")
    assert paras[2] == "Sixth keeps its text (Jones 2019)."
    assert stats3.total_markers == 3


def test_author_date_output_never_aliases_a_later_marker(tmp_path):
    """An inserted author-date citation can read exactly like a later
    author-suggested marker; each marker is still replaced at its own offset."""
    project = build_project(tmp_path, ["First (REF) and second (Smith et al., 2020)."])
    project.settings.citation_style = CitationStyle.APA
    (s,) = project.sentences
    smith = CitationCandidate(pmid="7", title="Smith paper", year=2020, journal="J",
                              authors=[Author(last_name="Smith", initials="A"),
                                       Author(last_name="Lee", initials="B"),
                                       Author(last_name="Kim", initials="C")])
    other = cand("8")
    project.evidence_map[s.id] = EvidenceRecord(
        sentence_id=s.id, selected=[smith, other], slot_sizes=[1, 1],
        review_decision=ReviewDecision.ACCEPTED)
    out = tmp_path / "out.docx"
    export_fresh(project, str(out))
    text = Document(str(out)).paragraphs[0].text
    assert text.startswith("First (Smith et al., 2020) and second (")
    assert "Author8" in text and text.count("(Smith et al., 2020)") == 1
