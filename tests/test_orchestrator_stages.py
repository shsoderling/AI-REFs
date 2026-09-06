"""The orchestrator runs search, verification and QA end to end on fakes."""

import pytest
from docx import Document

import src.pipeline.orchestrator as orch
from src.models.evidence import ConfidenceLevel, Verdict, VerificationStatus
from src.models.project import PipelineStage, ProjectState
from src.pipeline.llm_citation_agent import LLMCitationAgent
from src.storage.cache_db import CacheDB
from tests.fakes import FakeAnthropic, FakeBioRxiv, FakeEuropePMC, FakePubMed, article, submit_msg, tool_use_msg, verdict_msg

ABSTRACT_A = "Rac1 activation increased dendritic spine density in hippocampal neurons."
ABSTRACT_B = "Loss of Rac1 blocked spine enlargement after LTP induction."


@pytest.fixture
def docx_path(tmp_path):
    doc = Document()
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Spines are dynamic structures. Rac1 activation increases spine density (REF) "
                      "and loss of Rac1 blocks spine enlargement (REF).")
    path = tmp_path / "in.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def offline(monkeypatch):
    """Everything the orchestrator would build from the network becomes a fake."""
    pubmed = FakePubMed(by_search={"rac1 spine density": (["11"], 1), "rac1 ltp": (["12"], 1)},
                        articles={"11": article("11", "Rac1 and spines", abstract=ABSTRACT_A),
                                  "12": article("12", "Rac1 and LTP", abstract=ABSTRACT_B)})
    state = {"pubmed": pubmed, "epmc": FakeEuropePMC(), "biorxiv": FakeBioRxiv(), "clients": []}
    monkeypatch.setattr(orch, "CacheDB", lambda *a, **k: CacheDB(":memory:"))
    monkeypatch.setattr(orch, "PubMedClient", lambda **k: pubmed)
    monkeypatch.setattr(orch, "EuropePMCClient", lambda **k: state["epmc"])
    monkeypatch.setattr(orch, "BioRxivClient", lambda **k: state["biorxiv"])

    def agent_factory(script):
        def make(**kwargs):
            client = FakeAnthropic(script)
            state["clients"].append(client)
            return LLMCitationAgent(client=client, **kwargs)
        monkeypatch.setattr(orch, "LLMCitationAgent", make)
    state["use_script"] = agent_factory
    return state


def _project(docx_path, **settings):
    project = ProjectState(input_docx_path=docx_path)
    project.settings.anthropic_api_key = "k"
    project.settings.ncbi_email = "x@y.z"
    project.settings.domain_inference = False
    project.settings.reference_library_enabled = False
    project.settings.search_biorxiv = False
    for k, v in settings.items():
        setattr(project.settings, k, v)
    return project


def _search_script():
    return [
        tool_use_msg(("search_pubmed", {"query": "rac1 spine density"})),
        tool_use_msg(("fetch_articles", {"pmids": ["11"]})),
        submit_msg([{"pmid": "11"}], confidence="MEDIUM", score=60),
        tool_use_msg(("search_pubmed", {"query": "rac1 ltp"})),
        tool_use_msg(("fetch_articles", {"pmids": ["12"]})),
        submit_msg([{"pmid": "12"}], confidence="MEDIUM", score=60),
    ]


def test_full_run_searches_verifies_and_reports_stages(docx_path, offline):
    offline["use_script"](_search_script() + [
        verdict_msg("supports", "Rac1 activation increased dendritic spine density", "direct"),
        verdict_msg("not_supported", "", "the paper is about loss of function"),
    ])
    progress = []
    logs = []
    project = _project(docx_path)
    result = orch.PipelineOrchestrator(project, progress_callback=lambda *a: progress.append(a),
                                       log_callback=lambda lvl, m: logs.append(m)).run()

    assert result.current_stage == PipelineStage.COMPLETE
    stage_names = [p[0] for p in progress]
    assert stage_names == ["Parse Document", "Locate Markers", "AI Citation Search",
                           "Verify Citations", "Global QA", "Complete"]
    assert ("Verify Citations", 3, 5) in progress

    [ev] = list(result.evidence_map.values())          # one sentence, two markers
    assert [c.pmid for c in ev.selected] == ["11", "12"]
    assert [v.verdict for v in ev.verdicts] == [Verdict.SUPPORTS, Verdict.NOT_SUPPORTED]
    assert ev.confidence_level == ConfidenceLevel.LOW
    assert ev.verification_status == VerificationStatus.MISMATCH
    assert any(w.code == "not_supported" for w in ev.warnings)

    client = offline["clients"][0]
    first_prompt = client.calls[0]["messages"][0]["content"]
    assert "Section: Results" in first_prompt and "Spines are dynamic structures." in first_prompt
    assert 'ending at this (REF) marker: "Rac1 activation increases spine density"' in first_prompt
    # the verifier saw the second marker's sub-claim for the second paper
    verifier_prompt = client.calls[-1]["messages"][0]["content"]
    assert 'marker: "and loss of Rac1 blocks spine enlargement"' in verifier_prompt
    assert any("Verify [" in m for m in logs)


def test_verification_off_still_runs_deterministic_checks(docx_path, offline):
    preprint = article("11", "Preprint", doi="10.1101/2024.01.01.1", source="biorxiv",
                       publication_types=["Preprint"], published_doi="10.1/j")
    offline["pubmed"].articles["11"] = preprint
    offline["epmc"].by_doi["10.1/j"] = article("99", "Journal version", abstract=ABSTRACT_A)
    offline["use_script"](_search_script())

    project = _project(docx_path, verify_citations=False)
    result = orch.PipelineOrchestrator(project).run()

    [ev] = list(result.evidence_map.values())
    assert [c.pmid for c in ev.selected] == ["99", "12"]        # swapped, no verifier calls
    assert ev.verdicts == []
    assert [w.code for w in ev.warnings] == ["preprint_swapped"]
    assert len(offline["clients"][0].calls) == 6
