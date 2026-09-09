"""Tests for the OOXML fixture builders used by the DOCX edge-case suite."""

import pytest
from docx import Document

from tests.fixture_builders import (
    DocBuilder, W_NS, run_texts,
)


def test_builder_field_is_invisible_to_paragraph_text(tmp_path):
    b = DocBuilder()
    p = b.paragraph("Neurons fire")
    b.add_field(p, code=' ADDIN AIREFS.CITE {"x":1} ', result="1,2", superscript=True)
    b.add_text(p, ". Next")
    path = b.save(tmp_path / "f.docx")
    doc = Document(str(path))
    para = doc.paragraphs[0]
    assert para.text == "Neurons fire1,2. Next"
    assert run_texts(para) == ["Neurons fire", "", "", "", "1,2", "", ". Next"]
    assert [r.font.superscript for r in para.runs][4] is True


def test_builder_hyperlink_text_is_in_paragraph_text_but_not_runs(tmp_path):
    b = DocBuilder()
    p = b.paragraph("See ")
    b.add_hyperlink(p, "https://example.org", "the site")
    b.add_text(p, " here (REF).")
    path = b.save(tmp_path / "h.docx")
    para = Document(str(path)).paragraphs[0]
    assert para.text == "See the site here (REF)."
    assert run_texts(para) == ["See ", " here (REF)."]


def test_builder_split_instr_text_and_cross_paragraph_field(tmp_path):
    b = DocBuilder()
    p1 = b.paragraph("")
    b.add_field(p1, code=' ADDIN AIREFS.BIBL {"v":1} ', result="1. Entry one",
                split_code_at=12, end_in_new_paragraph=True, extra_paragraph_texts=["2. Entry two"])
    path = b.save(tmp_path / "s.docx")
    doc = Document(str(path))
    xml = doc.element.body.xml
    assert xml.count("<w:instrText") == 2
    assert doc.paragraphs[0].text == "1. Entry one"
    assert doc.paragraphs[1].text == "2. Entry two"


def test_builder_tracked_change_wrappers_and_multi_wt(tmp_path):
    b = DocBuilder()
    p = b.paragraph("A")
    b.add_text(p, "B", wrap="ins")
    b.add_text(p, "C", wrap="del")
    b.add_multi_wt_run(p, ["D", "E"])
    path = b.save(tmp_path / "t.docx")
    para = Document(str(path)).paragraphs[0]
    # python-docx ignores w:ins/w:del content in paragraph.text
    assert para.text == "ADE"
    body_xml = para._p.xml
    assert "<w:ins" in body_xml and "<w:del" in body_xml
    # make_run preserves whitespace on w:delText exactly as it does on w:t
    assert '<w:delText xml:space="preserve">C</w:delText>' in body_xml


def test_builder_tracked_move_wrappers(tmp_path):
    b = DocBuilder()
    p = b.paragraph("A")
    b.add_text(p, "B", wrap="moveFrom")
    b.add_text(p, "C", wrap="moveTo")
    path = b.save(tmp_path / "m.docx")
    para = Document(str(path)).paragraphs[0]
    # python-docx ignores w:moveFrom/w:moveTo content in paragraph.text
    assert para.text == "A"
    body_xml = para._p.xml
    assert "<w:moveFrom " in body_xml and "<w:moveTo " in body_xml
    # unlike w:del, moved-from text keeps w:t (ECMA-376 17.13.5.22)
    assert "<w:delText" not in body_xml
    assert '<w:t xml:space="preserve">B</w:t>' in body_xml


def test_builder_rejects_unknown_tracked_change_kind():
    b = DocBuilder()
    p = b.paragraph("A")
    with pytest.raises(ValueError):
        b.add_text(p, "B", wrap="moved")
