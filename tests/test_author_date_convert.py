"""Tests for numeric → author-date conversion in insert mode."""

from docx import Document

from src.models.citation import Author, CitationCandidate
from src.models.existing_refs import ExistingBibEntry, ExistingCitationMap, InTextCitation
from src.models.project import CitationStyle
from src.pipeline.author_date_convert import (
    build_author_date_bibliography, build_author_date_labels,
    convert_in_text_to_author_date,
)
from src.pipeline.renumbering import NewMarkerInfo, compute_renumbering
from src.services.docx_io import DocxHandler


def entry(num, body, matched=None, pmid="", doi=""):
    return ExistingBibEntry(original_number=num, raw_text=f"{num}. {body}",
                            body=body, matched_candidate=matched,
                            pmid=pmid, doi=doi)


def candidate(last_names, year, title="T"):
    return CitationCandidate(
        title=title, year=year,
        authors=[Author(last_name=ln, initials="X") for ln in last_names],
    )


class TestBuildLabels:
    def test_label_from_matched_candidate(self):
        e = entry(1, "ignored", matched=candidate(["Smith", "Doe", "Roe"], 2020))
        existing = ExistingCitationMap(bib_entries={1: e},
                                       references_heading_para_idx=5)
        result = build_author_date_labels(existing)
        assert result.labels[1] == "Smith et al., 2020"
        assert not result.missing

    def test_label_parsed_from_body_multi_author(self):
        e = entry(2, "Smith J, Doe A, Roe B. A title. J. 2019;1:1-2.")
        existing = ExistingCitationMap(bib_entries={2: e},
                                       references_heading_para_idx=5)
        result = build_author_date_labels(existing)
        assert result.labels[2] == "Smith et al., 2019"

    def test_label_parsed_two_authors(self):
        e = entry(3, "Smith J, Doe A. A title. J. 2018;1:1-2.")
        existing = ExistingCitationMap(bib_entries={3: e},
                                       references_heading_para_idx=5)
        result = build_author_date_labels(existing)
        assert result.labels[3] == "Smith & Doe, 2018"

    def test_label_single_author(self):
        e = entry(4, "Smith J. A title. J. 2017;1:1-2.")
        existing = ExistingCitationMap(bib_entries={4: e},
                                       references_heading_para_idx=5)
        result = build_author_date_labels(existing)
        assert result.labels[4] == "Smith, 2017"

    def test_unlabelable_entry_reported_missing(self):
        e = entry(5, "no year or author structure here")
        existing = ExistingCitationMap(bib_entries={5: e},
                                       references_heading_para_idx=5)
        result = build_author_date_labels(existing)
        assert result.missing == [5]


def _doc(tmp_path, body_builder):
    doc = Document()
    body_builder(doc)
    doc.add_paragraph("References")
    path = tmp_path / "d.docx"
    doc.save(str(path))
    return DocxHandler(str(path))


def _existing(entries, refs_idx):
    return ExistingCitationMap(bib_entries=entries,
                               references_heading_para_idx=refs_idx)


class TestConvertInText:
    def test_bracket_group_converted(self, tmp_path):
        handler = _doc(tmp_path, lambda d: d.add_paragraph(
            "Shown before [1, 2] and later [3]."))
        existing = _existing({
            1: entry(1, "x"), 2: entry(2, "x"), 3: entry(3, "x"),
        }, refs_idx=1)
        labels = {1: "Smith et al., 2020", 2: "Doe & Roe, 2019", 3: "Lee, 2021"}
        n = convert_in_text_to_author_date(handler, existing, labels)
        assert n == 2
        assert handler.get_paragraphs()[0].text == (
            "Shown before (Smith et al., 2020; Doe & Roe, 2019) and later (Lee, 2021).")

    def test_superscript_converted_and_unsuperscripted(self, tmp_path):
        def build(d):
            p = d.add_paragraph()
            p.add_run("As shown")
            r = p.add_run("1,2")
            r.font.superscript = True
            p.add_run(".")
        handler = _doc(tmp_path, build)
        existing = _existing({1: entry(1, "x"), 2: entry(2, "x")}, refs_idx=1)
        labels = {1: "Smith et al., 2020", 2: "Doe, 2019"}
        n = convert_in_text_to_author_date(handler, existing, labels)
        assert n == 1
        para = handler.get_paragraphs()[0]
        assert para.text == "As shown(Smith et al., 2020; Doe, 2019)."
        assert not any(r.font.superscript for r in para.runs)

    def test_non_citation_superscript_untouched(self, tmp_path):
        def build(d):
            p = d.add_paragraph()
            p.add_run("Ca")
            r = p.add_run("2+")
            r.font.superscript = True
        handler = _doc(tmp_path, build)
        existing = _existing({2: entry(2, "x")}, refs_idx=1)
        n = convert_in_text_to_author_date(handler, existing, {2: "Doe, 2019"})
        assert n == 0
        assert handler.get_paragraphs()[0].runs[1].text == "2+"

    def test_references_section_untouched(self, tmp_path):
        handler = _doc(tmp_path, lambda d: d.add_paragraph("Cite [1]."))
        # Add a paragraph after References containing [1]
        handler.doc.add_paragraph("1. Entry [1] mention.")
        existing = _existing({1: entry(1, "x")}, refs_idx=1)
        convert_in_text_to_author_date(handler, existing, {1: "Smith, 2020"})
        assert handler.get_paragraphs()[0].text == "Cite (Smith, 2020)."
        assert "[1]" in handler.get_paragraphs()[2].text


class TestAuthorDateBibliography:
    def test_sorted_unnumbered_merge(self):
        e1 = entry(1, "Zeta Z. Last alphabetically. J. 2020.")
        e2 = entry(2, "Alpha A. First alphabetically. J. 2019.")
        existing = ExistingCitationMap(
            bib_entries={1: e1, 2: e2},
            references_heading_para_idx=5,
            in_text_citations={0: [
                InTextCitation(char_offset=0, number=1),
                InTextCitation(char_offset=5, number=2),
            ]},
        )
        new_cand = CitationCandidate(
            pmid="9", title="Middle paper", year=2021,
            authors=[Author(last_name="Mmm", initials="M")],
            journal="J New",
        )
        result = compute_renumbering(existing, [
            NewMarkerInfo(para_index=0, char_offset=9, citations=[new_cand]),
        ])
        bib = build_author_date_bibliography(existing, result, CitationStyle.APA)
        assert len(bib) == 3
        assert bib[0].startswith("Alpha A.")
        assert bib[1].startswith("Mmm M.")
        assert bib[2].startswith("Zeta Z.")
        # Unnumbered
        assert not any(b[0].isdigit() for b in bib)
