"""Tests for the complex-field reader: fields are walked in document order,
split/cross-paragraph field codes are concatenated, nesting is recorded with
depth, and FieldIndex answers run-membership and result-span questions."""

from docx import Document

from src.services.docx_fields import (
    FieldIndex, classify_code, iter_complex_fields,
)
from tests.fixture_builders import DocBuilder, make_run


def _doc(tmp_path, build):
    b = DocBuilder()
    build(b)
    return Document(str(b.save(tmp_path / "d.docx")))


def test_single_field_code_and_result(tmp_path):
    def build(b):
        p = b.paragraph("x")
        b.add_field(p, code=' ADDIN AIREFS.CITE {"v":1} ', result="1", superscript=True)
    doc = _doc(tmp_path, build)
    fields = iter_complex_fields(doc.element.body)
    assert len(fields) == 1
    f = fields[0]
    assert f.code == 'ADDIN AIREFS.CITE {"v":1}'
    assert f.result_text == "1"
    assert f.complete and f.depth == 0 and f.kind == "airefs_cite"


def test_split_instr_text_is_concatenated(tmp_path):
    def build(b):
        p = b.paragraph("")
        b.add_field(p, code=' ADDIN AIREFS.CITE {"title":"a~b"} ', result="2", split_code_at=20)
    f = iter_complex_fields(_doc(tmp_path, build).element.body)[0]
    assert f.code == 'ADDIN AIREFS.CITE {"title":"a~b"}'


def test_field_spanning_paragraphs(tmp_path):
    def build(b):
        p = b.paragraph("")
        b.add_field(p, code=' ADDIN AIREFS.BIBL {"v":1} ', result="1. One",
                    end_in_new_paragraph=True, extra_paragraph_texts=["2. Two", "3. Three"])
    doc = _doc(tmp_path, build)
    f = iter_complex_fields(doc.element.body)[0]
    assert f.kind == "airefs_bibl" and f.complete
    assert len(f.paragraphs) == 3
    assert f.result_text == "1. One\n2. Two\n3. Three"


def test_nested_field_depth_and_foreign_classification(tmp_path):
    def build(b):
        p = b.paragraph("")
        b.add_field(p, code=" ADDIN EN.CITE <EndNote/> ", result="(Smith 2020)",
                    nested_code=" ADDIN EN.CITE.DATA ")
    fields = iter_complex_fields(_doc(tmp_path, build).element.body)
    assert [(f.kind, f.depth) for f in fields] == [("foreign", 0), ("foreign", 1)]


def test_fld_simple_and_unmatched_begin(tmp_path):
    def build(b):
        p = b.paragraph("")
        b.add_simple_field(p, instr=" PAGE ", result="7")
        p2 = b.paragraph("")
        p2._p.append(make_run(fldchar="begin"))
        p2._p.append(make_run(instr=" ADDIN AIREFS.CITE {} "))
    fields = iter_complex_fields(_doc(tmp_path, build).element.body)
    assert fields[0].kind == "other" and fields[0].result_text == "7"
    assert fields[1].complete is False


def test_tracked_change_wrappers_and_text_box_dedup(tmp_path):
    def build(b):
        p = b.paragraph("")
        b.add_field(p, code=' ADDIN AIREFS.CITE {"a":1} ', result="1", wrap="del")
        b.add_field(p, code=' ADDIN AIREFS.CITE {"a":2} ', result="2", wrap="ins")

        def fill(inner_p):
            for r in (make_run(fldchar="begin"), make_run(instr=' ADDIN AIREFS.CITE {"a":3} '),
                      make_run(fldchar="separate"), make_run("3"), make_run(fldchar="end")):
                inner_p.append(r)
        b.add_text_box(p, fill)
    doc = _doc(tmp_path, build)
    fields = iter_complex_fields(doc.element.body)
    assert len(fields) == 3
    assert fields[0].deleted is True and fields[1].deleted is False
    idx = FieldIndex(doc)
    assert idx.pending_tracked_changes is True
    assert idx.airefs_cite == 3


def test_tracked_change_inside_result_only_is_pending(tmp_path):
    """A field whose begin is clean but whose result run sits under w:ins is
    neither deleted nor inserted as a whole, yet still has a pending change."""
    def build(b):
        p = b.paragraph("")
        for r in (make_run(fldchar="begin"), make_run(instr=' ADDIN AIREFS.CITE {"a":5} '),
                  make_run(fldchar="separate")):
            p._p.append(r)
        b.add_text(p, "5", wrap="ins")
        p._p.append(make_run(fldchar="end"))
    doc = _doc(tmp_path, build)
    f = iter_complex_fields(doc.element.body)[0]
    assert f.complete and f.result_text == "5"
    assert f.deleted is False and f.inserted is False
    assert FieldIndex(doc).pending_tracked_changes is True


def test_field_index_roles_and_result_spans(tmp_path):
    def build(b):
        p = b.paragraph("Cells fire")
        b.add_field(p, code=' ADDIN AIREFS.CITE {"a":1} ', result="4,5", superscript=True)
        b.add_text(p, " and [2] more.")
    doc = _doc(tmp_path, build)
    para = doc.paragraphs[0]
    idx = FieldIndex(doc)
    roles = [idx.role(r) for r in para.runs]
    assert roles == [None, "marker", "code", "marker", "result", "marker", None]
    assert idx.in_field(para.runs[4]) and not idx.in_field(para.runs[0])
    assert idx.result_spans(para) == [(10, 13)]
    assert para.text[10:13] == "4,5"


def test_fields_in_tables_are_counted(tmp_path):
    def build(b):
        cell_p = b.add_table_cell_paragraph("Fig. 1 ")
        b.add_field(cell_p, code=' ADDIN AIREFS.CITE {"a":9} ', result="9")
    idx = FieldIndex(_doc(tmp_path, build))
    assert idx.airefs_cite == 1 and idx.in_tables == 1


def test_classify_code():
    assert classify_code(' ADDIN AIREFS.CITE {} ') == "airefs_cite"
    assert classify_code('ADDIN AIREFS.BIBL {}') == "airefs_bibl"
    assert classify_code(' ADDIN ZOTERO_ITEM CSL_CITATION {} ') == "foreign"
    assert classify_code(' ADDIN EN.CITE <x/> ') == "foreign"
    assert classify_code(' ADDIN Mendeley Bibliography CSL_BIBLIOGRAPHY ') == "foreign"
    assert classify_code(' PAGE ') == "other"
    assert classify_code(' HYPERLINK "x" ') == "other"
